#!/usr/bin/env python3
"""Initialize the paper trading Excel file with 3 tabs."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

wb = openpyxl.Workbook()

# Colors
green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF", size=11)
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'))

def style_header(ws, cols):
    for c in range(1, cols+1):
        cell = ws.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

# ============= TAB 1: Trade Log =============
ws1 = wb.active
ws1.title = "Trade Log"
headers1 = ["Date", "Stock", "Action", "Entry Price", "Exit Price", "Qty", "P&L (₹)", "Strategy", "Holding Days", "Notes"]
for i, h in enumerate(headers1, 1):
    ws1.cell(row=1, column=i, value=h)
style_header(ws1, len(headers1))
ws1.column_dimensions['A'].width = 12
ws1.column_dimensions['B'].width = 14
ws1.column_dimensions['C'].width = 10
ws1.column_dimensions['D'].width = 14
ws1.column_dimensions['E'].width = 14
ws1.column_dimensions['F'].width = 8
ws1.column_dimensions['G'].width = 14
ws1.column_dimensions['H'].width = 22
ws1.column_dimensions['I'].width = 12
ws1.column_dimensions['J'].width = 20

# ============= TAB 2: Strategy Performance =============
ws2 = wb.create_sheet("Strategy Performance")
strategies = ["Supertrend (10,3)", "EMA Crossover (6/30)", "RSI + MACD", "BB Squeeze", "VWAP Pullback"]
headers2 = ["Strategy", "Total Trades", "Wins", "Losses", "Win Rate %", "Total P&L (₹)", "Avg RR", "Max DD %"]
for i, h in enumerate(headers2, 1):
    ws2.cell(row=1, column=i, value=h)
style_header(ws2, len(headers2))
ws2.column_dimensions['A'].width = 22
for c in range(2, len(headers2)+1):
    ws2.column_dimensions[get_column_letter(c)].width = 14

for idx, s in enumerate(strategies):
    row = idx + 2
    ws2.cell(row=row, column=1, value=s)
    ws2.cell(row=row, column=2, value=0)
    ws2.cell(row=row, column=3, value=0)
    ws2.cell(row=row, column=4, value=0)
    ws2.cell(row=row, column=5, value=0)
    ws2.cell(row=row, column=6, value=0)
    ws2.cell(row=row, column=7, value=0)
    ws2.cell(row=row, column=8, value=0)
    for c in range(1, len(headers2)+1):
        ws2.cell(row=row, column=c).border = thin_border

# ============= TAB 3: Account Overview =============
ws3 = wb.create_sheet("Account Overview")

# Capital section
ws3.cell(row=1, column=1, value="PAPER TRADING ACCOUNT").font = Font(bold=True, size=14, color="4472C4")
ws3.merge_cells('A1:D1')

labels = [
    ("Starting Capital", "₹1,00,000"),
    ("Current Balance", "₹1,00,000"),
    ("Total P&L", "₹0"),
    ("Open Positions", "0"),
    ("Total Closed Trades", "0"),
]
for i, (lbl, val) in enumerate(labels):
    row = i + 3
    ws3.cell(row=row, column=1, value=lbl).font = Font(bold=True)
    ws3.cell(row=row, column=2, value=val)
    ws3.cell(row=row, column=2).font = Font(size=12)
    ws3.cell(row=row, column=1).border = thin_border
    ws3.cell(row=row, column=2).border = thin_border

ws3.cell(row=10, column=1, value="MONTHLY P&L").font = Font(bold=True, size=12, color="4472C4")
ws3.merge_cells('A10:D10')

# Month columns
months_headers = ["Month", "Starting Balance", "Ending Balance", "Monthly P&L", "Cumulative P&L"]
for i, h in enumerate(months_headers, 1):
    ws3.cell(row=11, column=i, value=h)
style_header(ws3, len(months_headers))

# Current month row
today = datetime.now()
current_month = today.strftime("%b %Y")
ws3.cell(row=12, column=1, value=current_month)
ws3.cell(row=12, column=2, value="₹1,00,000")
ws3.cell(row=12, column=3, value="₹1,00,000")
ws3.cell(row=12, column=4, value="₹0")
ws3.cell(row=12, column=5, value="₹0")
for c in range(1, 6):
    ws3.cell(row=12, column=c).border = thin_border

ws3.column_dimensions['A'].width = 22
ws3.column_dimensions['B'].width = 18
ws3.column_dimensions['C'].width = 18
ws3.column_dimensions['D'].width = 16
ws3.column_dimensions['E'].width = 18

# Save
path = "/home/ubuntu/Trading-BOT/trading/paper_trading/paper_trading.xlsx"
import os
os.makedirs(os.path.dirname(path), exist_ok=True)
wb.save(path)
print(f"✅ Excel created: {path}")
print(f"   Tab 1: Trade Log (trades with P&L)")
print(f"   Tab 2: Strategy Performance (per-strategy stats)")
print(f"   Tab 3: Account Overview (capital + monthly P&L)")
print(f"   Starting capital: ₹1,00,000")