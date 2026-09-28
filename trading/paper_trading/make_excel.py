#!/usr/bin/env python3
import openpyxl
from openpyxl.styles import Font
B="/home/ubuntu/Trading-BOT/trading/paper_trading"
E=f"{B}/paper_trading.xlsx"
wb=openpyxl.Workbook()
wb.remove(wb.active)
w1=wb.create_sheet("Trade Log")
for i,h in enumerate(["Date","Time","Stock","Strategy","Action","Entry","Exit","Qty","P&L","Exit Reason"],1): w1.cell(1,i,h).font=Font(bold=True)
w2=wb.create_sheet("Strategy Performance")
for i,h in enumerate(["Strategy","Trades","Wins","Losses","Win%","P&L","Avg RR","Now","Return%"],1): w2.cell(1,i,h).font=Font(bold=True)
w3=wb.create_sheet("Account Overview")
I=500000; w3.cell(1,1,"ACCOUNT OVERVIEW").font=Font(bold=True,size=14)
for i,(k,v) in enumerate([("Start",f"\u20b9{I:,}"),("Capital",f"\u20b9{I:,}"),("P&L","\u20b90"),("Trades",0),("Open",0),("Return","+0.00%")],2):
 w3.cell(i,1,k).font=Font(bold=True); w3.cell(i,2,v)
w3.cell(9,1,"50 STRATS | 100 STOCKS | SL 1.5x TP 3x | 10MIN").font=Font(bold=True,size=11)
wb.save(E); print(f"Excel created: {E}")