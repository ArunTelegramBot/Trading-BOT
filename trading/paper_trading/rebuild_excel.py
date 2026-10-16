#!/usr/bin/env python3
"""Full 3-tab rebuild of paper_trading.xlsx from state.json (redesigned template)."""
import json, os, warnings, sys
from datetime import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from openpyxl.formatting.rule import CellIsRule

warnings.filterwarnings('ignore')
B = "/home/ubuntu/Trading-BOT/trading/paper_trading"
E = f"{B}/paper_trading.xlsx"
S = f"{B}/state.json"

I = 500000
C = I // 50

ST = [
"S7/2","S10/2","S10/3","S14/3","S20/4",
"E5/13","E6/30","E8/21","E10/50","E20/50",
"RSI30","RSI25","RSI75","RSIDiv",
"MHist","MACross","MACDV",
"BBSqz","BB%<.2","BB%>.8","BBNar","BBExt",
"ADX25","ADX30","ADX20",
"D20","D10","D55",
"VSurge","VClimax","VAccum","VDist",
"KBreak","KCTouch","KReject",
"StrBreak","InBar","Engulf","3Bar",
"ATR1.5","ATR2","ATRAve",
"H9/18","H9/50",
"VWAPPull","VWAPDiv",
"ROCS","ROCV","PRIROC",
"Confluence"
]
IDLE = {"KCTouch","KReject","InBar","ATRAve","RSI30","RSIDiv"}

state = json.load(open(S))
trades = state.get("closed_trades", [])
strat = state.get("strategies", {})
psych = state.get("psych", {})

# ── Styles ──────────────────────────────────────
HDR = PatternFill("solid", fgColor="1E293B")
GRN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
SEC = PatternFill("solid", fgColor="EEF2F7")
MONTH = PatternFill("solid", fgColor="F8FAFC")
TB = Border(left=Side('thin', color='D1D5DB'), right=Side('thin', color='D1D5DB'),
            top=Side('thin', color='D1D5DB'), bottom=Side('thin', color='D1D5DB'))
WHITE = Font(color="FFFFFF", bold=True)
TITLE = Font(color="1E293B", size=18, bold=True)
SUBT = Font(color="64748B", size=10)
SECH = Font(color="334155", size=10, bold=True)
LABEL = Font(color="475569", size=11)
VALUE = Font(color="0F172A", size=11)
RUNCAP = Font(color="0F172A", size=11, bold=True)
BIG = Font(color="0F172A", size=14, bold=True)
GRAY = Font(color="94A3B8")
BLUE = Font(color="0000FF", bold=True)
NOTE = Font(color="64748B", size=9)
LOGSTK = Font(color="0F172A", bold=True)
LOGP = Font(color="475569")

wb = openpyxl.Workbook()
wb.remove(wb.active)

# ═══ Tab 1: Account Overview ═══════════════════
w3 = wb.create_sheet("Account Overview")
w3.cell(1, 1, "Account Overview").font = TITLE
w3.cell(2, 1, "Paper trading · all amounts in ₹").font = SUBT

def sec(row, txt):
    w3.cell(row, 1, txt).font = SECH
    for c in range(1, 3):
        w3.cell(row, c).fill = SEC

sec(3, "ACCOUNT")
w3.cell(4, 1, "Starting capital").font = LABEL
w3.cell(4, 2, I).font = VALUE
w3.cell(5, 1, "Running capital").font = LABEL
tc = sum(strat[s]['capital'] for s in ST) if all(s in strat for s in ST) else sum(s['capital'] for s in strat.values())
w3.cell(5, 2, round(tc, 0)).font = RUNCAP
w3.cell(6, 1, "Total P&L").font = LABEL
w3.cell(6, 2, "=B5-B4").font = BIG
w3.cell(7, 1, "Return").font = LABEL
w3.cell(7, 2, "=B6/B4").font = BIG

sec(9, "PERFORMANCE")
w3.cell(10, 1, "Closed trades").font = LABEL
w3.cell(10, 2, "=B11+B13").font = VALUE
wins = sum(1 for t in trades if t['pnl'] >= 0)
losses = sum(1 for t in trades if t['pnl'] < 0)
win_amt = sum(t['pnl'] for t in trades if t['pnl'] >= 0)
loss_amt = sum(t['pnl'] for t in trades if t['pnl'] < 0)
w3.cell(11, 1, "Winning trades").font = LABEL
w3.cell(11, 2, wins).font = VALUE
w3.cell(12, 1, "Winning P&L").font = LABEL
w3.cell(12, 2, round(win_amt, 0)).font = VALUE
w3.cell(13, 1, "Losing trades").font = LABEL
w3.cell(13, 2, losses).font = VALUE
w3.cell(14, 1, "Losing P&L").font = LABEL
w3.cell(14, 2, round(loss_amt, 0)).font = VALUE
w3.cell(15, 1, "Win rate").font = LABEL
w3.cell(15, 2, "=B11/B10").font = VALUE
w3.cell(16, 1, "Max drawdown").font = LABEL
w3.cell(16, 2, round(psych.get("dd", 0) / 100, 4)).font = VALUE
w3.cell(17, 1, "Loss streak").font = LABEL
w3.cell(17, 2, psych.get("loss_streak", 0)).font = VALUE

sec(19, "MONTHLY")
from collections import defaultdict
mdata = defaultdict(lambda: {"trades": 0, "wins": 0, "losses": 0, "pnl": 0})
for t in trades:
    d = str(t.get('exit_date', ''))[:7]
    if d:
        mdata[d]["trades"] += 1
        mdata[d]["pnl"] += t['pnl']
        if t['pnl'] >= 0: mdata[d]["wins"] += 1
        else: mdata[d]["losses"] += 1
ms = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
r = 20
for mkey in sorted(mdata.keys()):
    ys, mon = mkey.split('-')
    label = f"{ms[int(mon)-1]} {ys}"
    md = mdata[mkey]
    w3.cell(r, 1, label).font = Font(color="0F172A", bold=True)
    w3.cell(r, 2, round(md['pnl'], 0)).font = VALUE
    w3.cell(r+1, 1, "Trades").font = LABEL
    w3.cell(r+1, 2, md['trades']).font = VALUE
    w3.cell(r+2, 1, "Wins").font = LABEL
    w3.cell(r+2, 2, md['wins']).font = VALUE
    w3.cell(r+3, 1, "Losses").font = LABEL
    w3.cell(r+3, 2, md['losses']).font = VALUE
    for rr in range(r, r+4):
        w3.cell(rr, 1).fill = MONTH
        w3.cell(rr, 2).fill = MONTH
    r += 4

sec(29, "RULES")
rules = ["Max 2% risk per trade  ✓", "Trend filter  ✓", "Cool-off after 3 losses  ✓",
         "Half size at 10% drawdown", "50 strategies · 100 stocks",
         "Stop loss 1.5× · Target 3.0×"]
for i, txt in enumerate(rules):
    w3.cell(30 + i, 1, txt).font = LABEL

green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
w3.conditional_formatting.add('B6:B7', CellIsRule(operator='greaterThan', formula=['0'], fill=green))
w3.conditional_formatting.add('B6:B7', CellIsRule(operator='lessThan', formula=['0'], fill=red))
w3.conditional_formatting.add('B20:B50', CellIsRule(operator='greaterThan', formula=['0'], fill=green))
w3.conditional_formatting.add('B20:B50', CellIsRule(operator='lessThan', formula=['0'], fill=red))
for c in range(1, 3):
    w3.column_dimensions['A'].width = 24
    w3.column_dimensions['B'].width = 14

# ═══ Tab 2: Strategy Performance ════════════════
w2 = wb.create_sheet("Strategy Performance")
hdr2 = ["Strategy", "Return", "P&L ₹", "Win %", "Trades", "Wins", "Losses", "Avg RR", "Now ₹"]
for c, h in enumerate(hdr2, 1):
    cell = w2.cell(1, c, h)
    cell.fill = HDR; cell.font = WHITE
for i, s in enumerate(ST, 2):
    td = [t for t in trades if t['strategy'] == s]
    w = sum(1 for t in td if t['pnl'] >= 0)
    l = sum(1 for t in td if t['pnl'] < 0)
    pn = sum(t['pnl'] for t in td)
    aw = sum(t['pnl'] for t in td if t['pnl'] >= 0) / w if w > 0 else 0
    al = abs(sum(t['pnl'] for t in td if t['pnl'] < 0)) / l if l > 0 else 1
    cap = strat[s]['capital'] if s in strat else C
    cr = cap + pn
    w2.cell(i, 1, s)
    w2.cell(i, 2, f"=I{i}/$E$54-1")
    w2.cell(i, 3, round(pn, 0))
    w2.cell(i, 4, f"=IF(E{i}=0,0,F{i}/E{i})")
    w2.cell(i, 5, len(td))
    w2.cell(i, 6, w)
    w2.cell(i, 7, l)
    w2.cell(i, 8, round(aw / al, 2) if al > 0 else 0)
    w2.cell(i, 9, round(cr, 0))
    for c in range(1, 10):
        w2.cell(i, c).border = TB
        if s in IDLE and len(td) == 0:
            w2.cell(i, c).font = GRAY
w2.cell(54, 1, "Start capital per strategy (₹)").font = NOTE
w2.cell(54, 5, C).font = BLUE
w2.cell(55, 1, "Return = Now ₹ ÷ start capital − 1. Start capital inferred from the original file (₹10,000 per strategy).").font = NOTE
widths2 = [14, 10, 10, 8, 8, 7, 8, 9, 11]
for c, wdt in enumerate(widths2, 1):
    w2.column_dimensions[chr(64 + c)].width = wdt

# ═══ Tab 3: Trade Log ══════════════════════════
w1 = wb.create_sheet("Trade Log")
hdr1 = ["Stock", "P&L ₹", "Reason", "Strategy", "Action", "Entry", "Exit", "Qty", "Date", "Time"]
for c, h in enumerate(hdr1, 1):
    cell = w1.cell(1, c, h)
    cell.fill = HDR; cell.font = WHITE
    cell.border = TB
widths1 = [12, 9, 7, 12, 7, 10, 10, 5, 11, 7]
for c, wdt in enumerate(widths1, 1):
    w1.column_dimensions[chr(64 + c)].width = wdt

for idx, t in enumerate(trades):
    r = idx + 2
    w1.cell(r, 1, t['stock']).font = LOGSTK
    p = w1.cell(r, 2, round(t['pnl'], 0))
    p.fill = GRN if t['pnl'] >= 0 else RED
    p.font = VALUE
    w1.cell(r, 3, t.get('exit_reason', '')).font = VALUE
    w1.cell(r, 4, t['strategy']).font = VALUE
    w1.cell(r, 5, "SELL" if t['action'] == 'BUY' else "COVER").font = LOGP
    w1.cell(r, 6, t['entry']).font = VALUE
    w1.cell(r, 7, t['exit']).font = VALUE
    w1.cell(r, 8, t['qty']).font = VALUE
    ed = t['exit_date']
    if isinstance(ed, str):
        try:
            ed = datetime.strptime(ed, "%Y-%m-%d").date()
        except ValueError:
            try:
                ed = datetime.strptime(ed, "%Y-%m-%d %H:%M:%S").date()
            except ValueError:
                ed = ed
    w1.cell(r, 9, ed).font = LOGP
    w1.cell(r, 10, t.get('exit_time', '')).font = LOGP
    for c in range(1, 11):
        w1.cell(r, c).border = TB

# order: Account Overview first
wb._sheets = [wb['Account Overview'], wb['Strategy Performance'], wb['Trade Log']]
wb.save(E)

# ── Verify ─────────────────────────────────────
wb2 = openpyxl.load_workbook(E)
assert wb2.sheetnames == ['Account Overview', 'Strategy Performance', 'Trade Log'], wb2.sheetnames
assert wb2['Trade Log']['A1'].fill.start_color.rgb == '001E293B', wb2['Trade Log']['A1'].fill.start_color.rgb
assert wb2['Trade Log']['B2'].fill.start_color.rgb in ('00C6EFCE', '00FFC7CE'), wb2['Trade Log']['B2'].fill.start_color.rgb
assert wb2['Account Overview']['B6'].value == '=B5-B4', wb2['Account Overview']['B6'].value
assert len(wb2['Account Overview'].conditional_formatting._cf_rules) >= 2
log_rows = wb2['Trade Log'].max_row - 1
assert log_rows == len(trades), f"log rows {log_rows} != {len(trades)}"
wins_tab2 = sum(wb2['Strategy Performance'].cell(r, 6).value or 0 for r in range(2, 52))
print(f"OK: {len(trades)} trades, log rows {log_rows}, tab2 trades sum {wins_tab2 + sum(wb2['Strategy Performance'].cell(r,7).value or 0 for r in range(2,52))}")
print(f"Running capital: {round(tc,0)}")
print("REBUILD DONE")
